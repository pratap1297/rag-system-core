"""
Enhanced Text Extractor for RAG System
"""

import os
import io
import mimetypes
from pathlib import Path
from typing import Dict, Any, Optional, List, Tuple
from dataclasses import dataclass
from dotenv import load_dotenv

from ..core.logging_system import get_logger
from ..core.exceptions import FileProcessingError, AzureServiceError
from ..core.error_handler import with_error_handling
from ..core.monitoring import get_performance_monitor
from .visio_utils import convert_visio_to_pdf
from .azure_ocr import azure_ocr_image
import io
from pdf2image import convert_from_path
from PIL import Image

try:
    import pytesseract
    TESSERACT_AVAILABLE = True
except ImportError:
    TESSERACT_AVAILABLE = False

load_dotenv()

@dataclass
class ExtractionResult:
    """Result of text extraction"""
    text: str
    metadata: Dict[str, Any]
    confidence: float
    extraction_method: str
    page_count: Optional[int] = None
    errors: List[str] = None

class TextExtractor:
    """Enhanced text extractor with multi-format support"""
    
    def __init__(self, config: Optional[Dict[str, Any]] = None):
        self.config = config or {}
        self.logger = get_logger("text_extractor")
        self.monitor = get_performance_monitor()
        self.use_tesseract_fallback = self.config.get('use_tesseract_fallback', True)
        self.use_direct_pdf_text = self.config.get('use_direct_pdf_text', False)
        self.logger.info("Text extractor initialized")
    
    @with_error_handling("text_extractor", "extract_text")
    def extract_text(self, file_path: str, **kwargs) -> ExtractionResult:
        """Extract text from document"""
        
        try:
            file_path = Path(file_path)
            
            if not file_path.exists():
                raise FileProcessingError(f"File not found: {file_path}")
            
            file_extension = file_path.suffix.lower()
            
            self.logger.info(f"Extracting text from {file_path}")
            
            # Visio support
            if file_extension in ['.vsd', '.vsdx']:
                pdf_path = convert_visio_to_pdf(file_path)
                return self._extract_from_pdf_with_ocr(pdf_path)
            
            # PDF: Use Azure OCR by default
            if file_extension == '.pdf':
                if self.use_direct_pdf_text:
                    result = self._extract_from_pdf(file_path)
                    if result.text.strip():
                        return result
                return self._extract_from_pdf_with_ocr(file_path)
            
            # Images: Use Azure OCR
            if file_extension in ['.png', '.jpg', '.jpeg']:
                return self._extract_from_image_with_ocr(file_path)
            
            # Route to appropriate extraction method
            elif file_extension in ['.docx', '.doc']:
                result = self._extract_from_word(file_path)
            elif file_extension == '.txt':
                result = self._extract_from_text(file_path)
            elif file_extension == '.xlsx':
                result = self._extract_from_excel(file_path)
            else:
                raise FileProcessingError(f"Unsupported file format: {file_extension}")
            
            self.logger.info(f"Text extraction completed: {len(result.text)} characters")
            return result
            
        except Exception as e:
            raise FileProcessingError(f"Failed to extract text: {str(e)}")
    
    def _extract_from_pdf(self, file_path: Path) -> ExtractionResult:
        """Extract text from PDF files"""
        
        try:
            import PyPDF2
            
            text_content = []
            metadata = {}
            
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                
                page_count = len(pdf_reader.pages)
                for page_num, page in enumerate(pdf_reader.pages):
                    try:
                        page_text = page.extract_text()
                        if page_text.strip():
                            text_content.append(f"[Page {page_num + 1}]\n{page_text}")
                    except Exception as e:
                        self.logger.warning(f"Failed to extract page {page_num + 1}: {e}")
                
                full_text = '\n\n'.join(text_content)
                
                return ExtractionResult(
                    text=full_text,
                    metadata=metadata,
                    confidence=0.9 if full_text.strip() else 0.1,
                    extraction_method="PyPDF2",
                    page_count=page_count
                )
                
        except ImportError:
            raise FileProcessingError("PyPDF2 not available")
        except Exception as e:
            raise FileProcessingError(f"PDF extraction failed: {e}")
    
    def _extract_from_word(self, file_path: Path) -> ExtractionResult:
        """Extract text from Word documents"""
        
        try:
            import docx
            
            doc = docx.Document(file_path)
            
            text_content = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)
            
            full_text = '\n\n'.join(text_content)
            
            return ExtractionResult(
                text=full_text,
                metadata={},
                confidence=0.95,
                extraction_method="python-docx"
            )
            
        except ImportError:
            raise FileProcessingError("python-docx not available")
        except Exception as e:
            raise FileProcessingError(f"Word extraction failed: {e}")
    
    def _extract_from_text(self, file_path: Path) -> ExtractionResult:
        """Extract text from plain text files"""
        
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
            
            return ExtractionResult(
                text=text,
                metadata={'encoding': 'utf-8'},
                confidence=1.0,
                extraction_method="direct_read"
            )
            
        except Exception as e:
            raise FileProcessingError(f"Text extraction failed: {e}")
    
    def _extract_from_excel(self, file_path: Path) -> ExtractionResult:
        """Extract text from Excel (.xlsx) files"""
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text_content = []
            for sheet in wb.worksheets:
                text_content.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    row_text = ' '.join(str(cell) for cell in row if cell is not None)
                    if row_text.strip():
                        text_content.append(row_text)
            full_text = '\n'.join(text_content)
            return ExtractionResult(
                text=full_text,
                metadata={'sheets': [s.title for s in wb.worksheets]},
                confidence=1.0 if full_text.strip() else 0.1,
                extraction_method="openpyxl"
            )
        except ImportError:
            raise FileProcessingError("openpyxl not available")
        except Exception as e:
            raise FileProcessingError(f"Excel extraction failed: {e}")
    
    def _extract_from_pdf_with_ocr(self, file_path: Path) -> ExtractionResult:
        try:
            images = convert_from_path(str(file_path))
            all_text = []
            for i, image in enumerate(images):
                buf = io.BytesIO()
                image.save(buf, format="JPEG")
                buf.seek(0)
                try:
                    text = azure_ocr_image(buf.read())
                except Exception as azure_exc:
                    if self.use_tesseract_fallback and TESSERACT_AVAILABLE:
                        text = pytesseract.image_to_string(image)
                        self.logger.warning(f"Azure OCR failed, used Tesseract fallback on page {i+1}: {azure_exc}")
                    else:
                        raise FileProcessingError(f"Azure OCR failed and no fallback: {azure_exc}")
                all_text.append(f"[Page {i+1}]\n{text}")
            full_text = "\n\n".join(all_text)
            return ExtractionResult(
                text=full_text,
                metadata={'pages': len(images)},
                confidence=1.0 if full_text.strip() else 0.1,
                extraction_method="azure_ocr"
            )
        except Exception as e:
            raise FileProcessingError(f"PDF OCR extraction failed: {e}")

    def _extract_from_image_with_ocr(self, file_path: Path) -> ExtractionResult:
        try:
            with open(file_path, 'rb') as f:
                image_bytes = f.read()
            try:
                text = azure_ocr_image(image_bytes)
            except Exception as azure_exc:
                if self.use_tesseract_fallback and TESSERACT_AVAILABLE:
                    image = Image.open(file_path)
                    text = pytesseract.image_to_string(image)
                    self.logger.warning(f"Azure OCR failed, used Tesseract fallback: {azure_exc}")
                else:
                    raise FileProcessingError(f"Azure OCR failed and no fallback: {azure_exc}")
            return ExtractionResult(
                text=text,
                metadata={'source': str(file_path)},
                confidence=1.0 if text.strip() else 0.1,
                extraction_method="azure_ocr"
            )
        except Exception as e:
            raise FileProcessingError(f"Image OCR extraction failed: {e}")

    def get_supported_formats(self) -> List[str]:
        """Get list of supported file formats"""
        return ['.txt', '.pdf', '.docx', '.doc', '.xlsx', '.png', '.jpg', '.jpeg']
    
    def validate_file(self, file_path: str) -> Tuple[bool, str]:
        """Validate if file can be processed"""
        
        try:
            file_path = Path(file_path)
            
            if not file_path.exists():
                return False, "File does not exist"
            
            if file_path.stat().st_size == 0:
                return False, "File is empty"
            
            file_extension = file_path.suffix.lower()
            if file_extension not in self.get_supported_formats():
                return False, f"Unsupported format: {file_extension}"
            
            return True, "File is valid"
            
        except Exception as e:
            return False, f"Validation error: {str(e)}" 